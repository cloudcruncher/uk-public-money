from __future__ import annotations

from dataclasses import replace

import pytest
import yaml
from openpyxl import Workbook

from conftest import ROOT
from ukpm import registry, sheets, units
from ukpm.parsers import normalise_fiscal_year
from ukpm.provenance import ProvenanceError, provenance_id, validate
from ukpm.units import MissingDenominator

# --- provenance gate ---


def test_complete_record_passes(prov):
    assert validate(prov) is prov


@pytest.mark.parametrize(
    "field,bad",
    [
        ("publisher", ""),
        ("publication_date", "16/07/2026"),
        ("source_url", "http://assets.publishing.service.gov.uk/x.xlsx"),
        ("source_url", "https://evil.example.com/x.xlsx"),
        ("landing_url", None),
        ("file_sha256", "abc"),
        ("sheet", None),  # no sheet and no api_query
        ("cell_or_range", None),  # sheet set without cell
        ("licence", "UNVERIFIED"),
        ("status", "estimate"),
        ("vintage", " "),
        ("retrieved_at", "yesterday"),
        ("transformations", None),
    ],
)
def test_incomplete_record_rejected(prov, field, bad):
    with pytest.raises(ProvenanceError):
        validate(replace(prov, **{field: bad}))


def test_api_query_can_replace_sheet(prov):
    validate(replace(prov, sheet=None, cell_or_range=None, api_query="series=MGRZ"))


def test_provenance_id_is_stable_and_content_sensitive(prov):
    assert provenance_id("pesa", prov) == provenance_id("pesa", prov)
    assert provenance_id("pesa", prov) != provenance_id("pesa", replace(prov, cell_or_range="F21"))
    assert provenance_id("pesa", prov).startswith("pesa-")


# --- registry ---


def test_real_registry_loads():
    s = registry.load(ROOT / "data" / "sources.yml")
    assert len(s) >= 4
    by_id = {x.id: x for x in s}
    assert by_id["pesa"].enabled and by_id["ons_psf"].enabled and not by_id["ons_lms"].enabled
    assert all(x.verdict in ("USABLE", "PARTIAL") for x in s if x.enabled)


def _write(tmp_path, entries):
    p = tmp_path / "s.yml"
    p.write_text(yaml.safe_dump(entries))
    return p


def _entry(**over):
    e = {
        "id": "x",
        "publisher": "p",
        "dataset": "d",
        "landing_url": "https://www.gov.uk/a",
        "licence": "OGL-UK-3.0",
        "verdict": "USABLE",
        "enabled": True,
        "files": [
            {
                "key": "k",
                "edition": "e",
                "publication_date": "2026-01-01",
                "file_url": "https://www.gov.uk/f.xlsx",
                "landing_url": "https://www.gov.uk/a",
                "format": "xlsx",
                "sha256": "b" * 64,
                "role": "r",
            }
        ],
    }
    e.update(over)
    return e


def test_registry_accepts_minimal(tmp_path):
    assert registry.load(_write(tmp_path, [_entry()]))[0].id == "x"


@pytest.mark.parametrize(
    "over",
    [
        {"verdict": "UNVERIFIED"},
        {"licence": "UNVERIFIED"},
        {"landing_url": "http://www.gov.uk/a"},
        {"landing_url": "https://example.com/a"},
        {"files": [{**_entry()["files"][0], "sha256": None}]},
        {"files": [{**_entry()["files"][0], "publication_date": "soon"}]},
        {"files": [_entry()["files"][0], _entry()["files"][0]]},
    ],
)
def test_registry_rejects(tmp_path, over):
    with pytest.raises(registry.RegistryError):
        registry.load(_write(tmp_path, [_entry(**over)]))


def test_registry_rejects_duplicate_id_and_missing_key(tmp_path):
    with pytest.raises(registry.RegistryError):
        registry.load(_write(tmp_path, [_entry(), _entry()]))
    e = _entry()
    del e["publisher"]
    with pytest.raises(registry.RegistryError):
        registry.load(_write(tmp_path, [e]))


def test_disabled_source_may_lack_url_and_sha(tmp_path):
    f = {**_entry()["files"][0], "file_url": None, "sha256": None}
    registry.load(_write(tmp_path, [_entry(enabled=False, verdict="UNVERIFIED", files=[f])]))


# --- units ---


def test_scale_appends_transformation(make_row):
    r = units.scale(make_row(value=4330.0, unit="gbp_m"), 0.001, "£m -> £bn (x 0.001)", "gbp_bn")
    assert r.value == pytest.approx(4.33) and r.unit == "gbp_bn"
    assert r.prov.transformations[-1].op == "scale"
    assert len(r.prov.transformations) == 2


def test_deflate_and_pct_gdp_round_trip(make_row):
    defl = [
        make_row("obr.gdp_deflator", "2015-16", 80.0, unit="index"),
        make_row("obr.gdp_deflator", "2025-26", 100.0, unit="index"),
    ]
    nominal = make_row(year="2015-16", value=800.0)
    real = units.deflate(nominal, defl, "2025-26")
    assert real.value == pytest.approx(1000.0)
    assert real.value * 80.0 / 100.0 == pytest.approx(nominal.value)
    assert real.prov.transformations[-1].op == "deflate"
    assert len(real.prov.transformations[-1].inputs) == 2

    gdp = [make_row("obr.gdp_nominal", "2015-16", 2000.0)]
    pct = units.pct_gdp(nominal, gdp)
    assert pct.value == pytest.approx(40.0) and pct.unit == "pct"
    assert pct.value / 100 * gdp[0].value == pytest.approx(nominal.value)
    validate(pct.prov)


def test_missing_denominator_and_unit_guard(make_row):
    with pytest.raises(MissingDenominator):
        units.pct_gdp(make_row(year="2015-16"), [])
    with pytest.raises(MissingDenominator):
        units.deflate(make_row(year="2015-16"), [], "2025-26")
    with pytest.raises(ValueError):
        units.pct_gdp(make_row(unit="count"), [])


def test_forecast_denominator_makes_result_forecast(make_row):
    gdp = [make_row("obr.gdp_nominal", "2027-28", 3000.0, status="forecast")]
    assert units.pct_gdp(make_row(year="2027-28"), gdp).prov.status == "forecast"


# --- helpers and sheets ---


def test_fiscal_year_normalisation():
    assert normalise_fiscal_year("2015/16") == "2015-16"
    assert normalise_fiscal_year("2015-2016") == "2015-16"
    assert normalise_fiscal_year("1999-00") == "1999-00"
    with pytest.raises(ValueError):
        normalise_fiscal_year("2015-18")


def test_a1():
    assert sheets.a1(0, 0) == "A1"
    assert sheets.a1(45, 1) == "B46"
    assert sheets.a1(0, 26) == "AA1"


def test_sheets_a1_mapping_matches_openpyxl(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws["C5"] = "Total managed expenditure"
    ws["F7"] = 1234.5
    ws["AB9"] = "far"
    path = tmp_path / "t.xlsx"
    wb.save(path)

    book = sheets.open_book(path)
    assert book.sheet_names == ["T"]
    r, c = book.find("T", "total managed expenditure")
    assert sheets.a1(r, c) == "C5"
    assert book.rows("T")[6][5] == 1234.5
    for addr in ("C5", "F7", "AB9"):
        cell = ws[addr]
        assert book.rows("T")[cell.row - 1][cell.column - 1] == cell.value
        assert sheets.a1(cell.row - 1, cell.column - 1) == addr
    with pytest.raises(LookupError):
        book.find("T", "absent")


# --- layman-and-sources units ---


def test_per_person_scales_billions_to_pounds(make_row):
    pop = make_row("ons_population.uk", "2024-25", 50_000_000.0, unit="persons")
    out = units.per_person(make_row("pesa.tme", "2024-25", 100.0), [pop])
    assert out.unit == "gbp_per_person" and out.value == pytest.approx(2000.0)
    assert out.prov.transformations[-1].inputs == (units._pid(pop),)


def test_per_person_without_population_year_raises(make_row):
    with pytest.raises(MissingDenominator):
        units.per_person(make_row("pesa.tme", "2024-25", 100.0), [])


def test_share_per_100_and_share_of_total(make_row):
    part, total = make_row("a", value=25.0), make_row("b", value=200.0)
    assert units.share(part, total).unit == "gbp_per_100"
    assert units.share(part, total).value == pytest.approx(12.5)
    assert units.share(part, total, "share_of_total").unit == "pct"
    with pytest.raises(ValueError):
        units.share(part, total, "nope")
    with pytest.raises(ValueError):
        units.share(part, make_row("b", value=1.0, edition="PESA 2025"))


def test_per_claimant(make_row):
    spend = make_row("s", value=2.0)
    case = make_row("c", value=500.0, unit="thousand_cases")
    assert units.per_claimant(spend, case).value == pytest.approx(4000.0)
    with pytest.raises(ValueError):
        units.per_claimant(spend, make_row("c", value=500.0))
    with pytest.raises(MissingDenominator):
        units.per_claimant(spend, make_row("c", value=0.0, unit="thousand_cases"))


def test_remainder_and_sum_rows(make_row):
    total = make_row("t", value=100.0)
    parts = [make_row("p1", value=30.0), make_row("p2", value=20.0)]
    assert units.remainder(total, parts, "other").value == pytest.approx(50.0)
    with pytest.raises(ValueError):
        units.remainder(total, [make_row("p", value=1.0, edition="X")], "other")
    assert units.remainder(total, [make_row("p", value=1.0, edition="X")], "o", strict=False)
    s = units.sum_rows(parts[0], parts[1:], "sum_months", "d")
    assert s.value == pytest.approx(50.0)
    assert s.prov.transformations[-1].inputs == (units._pid(parts[1]),)


def test_registry_verdict_reason_is_optional_and_loaded():
    srcs = registry.load(ROOT / "data/sources.yml")
    by = {s.id: s for s in srcs.values()} if isinstance(srcs, dict) else {s.id: s for s in srcs}
    assert by["ons_psf"].verdict == "USABLE"
    assert by["obr_efo"].verdict_reason
    assert by["ons_population"].licence == "OGL-UK-3.0"


def test_per_household_scales_and_cites_household_count(make_row):
    hh = make_row("ons_households.uk", "2025-26", 29_003.0, unit="thousands")
    out = units.per_household(make_row("pesa.tme", "2025-26", 29.003), [hh])
    assert out.unit == "gbp_per_household" and out.value == pytest.approx(1000.0)
    t = out.prov.transformations[-1]
    assert t.op == "per_household" and t.inputs == (units._pid(hh),)
    assert "April-June" in t.detail and "fiscal year" in t.detail


def test_per_household_without_household_year_raises(make_row):
    with pytest.raises(MissingDenominator):
        units.per_household(make_row("pesa.tme", "2025-26", 100.0), [])


def test_per_household_rejects_wrong_unit(make_row):
    hh = make_row("ons_households.uk", "2025-26", 29_003_000.0, unit="persons")
    with pytest.raises(ValueError):
        units.per_household(make_row("pesa.tme", "2025-26", 100.0), [hh])
