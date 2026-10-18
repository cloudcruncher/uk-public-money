"""MHCLG local authority revenue outturn (England, Table 1a) against a trimmed real fixture."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, mhclg_la_outturn
from ukpm.provenance import problems
from ukpm.registry import FileEntry, Source

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests/fixtures/mhclg/table_1a.trimmed.xlsx"
SHA = "2807508a4030dedffd0c3377e89dd829b40aada18739ba2cf245a12288b23488"
URL = "https://assets.publishing.service.gov.uk/media/6aa91ca70420ac660805f7d9/RO_2025-26_Table_1a_1b.ods"
LANDING = "https://www.gov.uk/government/statistics/x"
SID = "mhclg_la_outturn"


def meta_for(path: Path) -> FileMeta:
    f = FileEntry(f"{SID}_k", f"{SID} edition", "2026-09-17", URL, LANDING, "ods", SHA, "main")
    s = Source(
        id=SID, publisher="Test publisher", dataset=f"{SID} dataset", landing_url=LANDING,
        licence="OGL-UK-3.0", licence_evidence=None, cadence=None, coverage={},
        verdict="USABLE", enabled=True, fixture_dir=None, user_agent=None, files=(f,),
    )  # fmt: skip
    return FileMeta(s, f, SHA, URL, "2026-10-18T00:00:00Z", path)


def run(path: Path = FIXTURE):
    return mhclg_la_outturn.parse(path, meta_for(path))


@pytest.fixture(scope="module")
def rows():
    return run()


def one(rows, slug, fy):
    (r,) = [x for x in rows if x.series_id == f"mhclg_la_outturn.{slug}" and x.fiscal_year == fy]
    return r


def edited(tmp_path, fn):
    wb = openpyxl.load_workbook(FIXTURE)
    fn(wb["Table_1a"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    return p


def test_provenance_complete_and_series_declared(rows):
    assert len(rows) == 12 * 5
    for r in rows:
        assert problems(r.prov) == [], (r.series_id, r.fiscal_year)
        assert (r.prov.file_sha256, r.prov.source_url) == (SHA, URL)
        assert r.prov.licence == "OGL-UK-3.0" and r.prov.status == "outturn"
        assert (r.unit, r.geography) == ("gbp_m", "England")
        assert r.series_id in mhclg_la_outturn.SERIES
    assert {r.series_id for r in rows} == set(mhclg_la_outturn.SERIES)


def test_cell_matches_workbook(rows):
    ws = openpyxl.load_workbook(FIXTURE, data_only=True)["Table_1a"]
    for r in rows:
        assert ws[r.prov.cell_or_range].value == pytest.approx(r.value), r.series_id


@pytest.mark.parametrize(
    ("slug", "year", "cell", "value"),
    [
        ("education", "2021-22", "B7", 34838.544),
        ("education", "2025-26", "F7", 49402.60738112),
        ("social_care", "2025-26", "F10", 44375.10174755),
        ("childrens_social_care", "2025-26", "F11", 16948.0980948),
        ("police", "2024-25", "E15", 16908.554),
        ("total", "2021-22", "B19", 107050.069),
        ("total", "2025-26", "F19", 144434.26457926),
    ],
)
def test_values(rows, slug, year, cell, value):
    r = one(rows, slug, year)
    assert r.prov.cell_or_range == cell
    assert r.prov.sheet == "Table_1a"
    assert r.value == pytest.approx(value)


def test_caveats_in_every_description():
    for s in mhclg_la_outturn.SERIES.values():
        assert "England only" in s["description"]
        assert "provisional" in s["description"]
        assert "Better Care Fund" in s["description"]
    assert "academies" in mhclg_la_outturn.SERIES["mhclg_la_outturn.education"]["description"]


def test_total_mismatch_is_layout_error(tmp_path):
    def f(ws):
        ws["F19"] = 150000

    with pytest.raises(LayoutError, match="do not add up"):
        run(edited(tmp_path, f))


def test_unit_change_is_layout_error(tmp_path):
    def f(ws):
        ws["A4"] = "Figures are reported in £ thousands."

    with pytest.raises(LayoutError, match="unit"):
        run(edited(tmp_path, f))


def test_title_change_is_layout_error(tmp_path):
    def f(ws):
        ws["A1"] = "Table 1a: Service expenditure in real terms"

    with pytest.raises(LayoutError, match="title"):
        run(edited(tmp_path, f))


def test_missing_service_is_layout_error(tmp_path):
    def f(ws):
        ws["A15"] = "Policing"

    with pytest.raises(LayoutError, match="Police"):
        run(edited(tmp_path, f))


def test_missing_sheet_is_layout_error(tmp_path):
    def f(ws):
        ws.title = "Other"

    with pytest.raises(LayoutError, match="Table_1a"):
        run(edited(tmp_path, f))


def test_non_consecutive_years_is_layout_error(tmp_path):
    def f(ws):
        ws["D6"] = "2024-25"

    with pytest.raises(LayoutError, match="consecutive"):
        run(edited(tmp_path, f))
