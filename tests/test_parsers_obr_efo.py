"""OBR EFO March 2026 expenditure parser against the unmodified workbook."""

from __future__ import annotations

from functools import cache
from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, obr_efo
from ukpm.provenance import problems
from ukpm.registry import FileEntry, Source

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "tests/fixtures/obr_efo/efo-march-2026-detailed-forecast-tables-expenditure.xlsx"
SHA = "700bfd700c872794e5fe8cf164086cd38a63850db26427b9d4b2e363d177d9a1"
URL = "https://obr.uk/docs/d055fbf02d5b3g6jq8l2/efo-march-2026-detailed-forecast-tables-expenditure.xlsx"
YEARS = ["2024-25", "2025-26", "2026-27", "2027-28", "2028-29", "2029-30", "2030-31"]
UC, OTHER, NI = (
    "obr_efo.welfare_cap.universal_credit",
    "obr_efo.welfare_cap.other_dwp",
    "obr_efo.welfare_cap.ni_social_security",
)


def _meta(path: Path = XLSX) -> FileMeta:
    f = FileEntry(
        "obr_efo_mar2026_expenditure",
        "OBR EFO March 2026",
        "2026-03-03",
        URL,
        "https://obr.uk/efo/economic-and-fiscal-outlook-march-2026/",
        "xlsx",
        SHA,
        "efo_expenditure",
    )
    s = Source(
        "obr_efo",
        "Office for Budget Responsibility",
        "Economic and fiscal outlook, detailed forecast tables (expenditure)",
        "https://obr.uk/efo/economic-and-fiscal-outlook-march-2026/",
        "OGL-UK-3.0",
        "test",
        None,
        {},
        "PARTIAL",
        True,
        None,
        None,
        (f,),
    )
    return FileMeta(s, f, SHA, URL, "2026-10-15T00:00:00Z", path)


@cache
def _rows():
    return obr_efo.parse(XLSX, _meta())


def _by(sid):
    return {r.fiscal_year: r for r in _rows() if r.series_id == sid}


def test_series_ids_match_series_dict():
    assert set(obr_efo.SERIES) == {UC, OTHER, NI}
    assert {r.series_id for r in _rows()} == set(obr_efo.SERIES)


def test_year_header_row_and_unit_in_source():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["4.9"]
    assert [ws.cell(5, c).value for c in range(3, 10)] == YEARS
    assert ws["C3"].value == "£ billion"
    assert ws["B18"].value == "Universal credit"
    assert ws["B34"].value == "NI social security in welfare cap"


@pytest.mark.parametrize(
    "sid,geo,row,vals",
    [
        (UC, "GB", 18, {"2024-25": 52.83499899882645, "2025-26": 66.41080001147837}),
        (OTHER, "GB", 19, {"2024-25": 2.0198741370396363, "2030-31": 1.9146847282577255}),
        (NI, "Northern Ireland", 34, {"2024-25": 5.7091400000000005, "2027-28": 6.473749}),
    ],
)
def test_values_cells_and_geography(sid, geo, row, vals):
    got = _by(sid)
    assert list(got) == YEARS
    for fy, v in vals.items():
        r = got[fy]
        assert r.value == v
        assert r.unit == "gbp_bn" and r.geography == geo and r.period_basis == "fiscal_year"
        assert r.prov.sheet == "4.9"
        assert r.prov.cell_or_range == f"{'CDEFGHI'[YEARS.index(fy)]}{row}"


def test_status_from_marker_row():
    for sid in obr_efo.SERIES:
        got = _by(sid)
        assert got["2024-25"].prov.status == "outturn"
        assert all(got[y].prov.status == "forecast" for y in YEARS[1:])


def test_provenance_complete_and_vintage():
    for r in _rows():
        assert problems(r.prov) == []
        assert r.prov.vintage == "OBR EFO March 2026"
        assert r.prov.file_sha256 == SHA and r.prov.source_url == URL


def _edit(tmp_path, fn):
    wb = openpyxl.load_workbook(XLSX)
    fn(wb)
    p = tmp_path / "edited.xlsx"
    wb.save(p)
    return p


def test_no_marker_row_means_all_forecast(tmp_path):
    def f(wb):
        ws = wb["4.9"]
        ws["C4"].value = None
        ws["D4"].value = None

    rows = obr_efo.parse(_edit(tmp_path, f), _meta())
    assert {r.prov.status for r in rows} == {"forecast"}


def test_missing_unit_raises(tmp_path):
    p = _edit(tmp_path, lambda wb: wb["4.9"].__setitem__("C3", "£ million"))
    with pytest.raises(LayoutError):
        obr_efo.parse(p, _meta())


def test_missing_sheet_raises(tmp_path):
    p = _edit(tmp_path, lambda wb: setattr(wb["4.9"], "title", "4.99"))
    with pytest.raises(LayoutError):
        obr_efo.parse(p, _meta())


def test_year_header_missing_raises(tmp_path):
    def f(wb):
        for c in "CDEFGHI":
            wb["4.9"][f"{c}5"].value = "x"

    with pytest.raises(LayoutError):
        obr_efo.parse(_edit(tmp_path, f), _meta())


def test_missing_label_raises(tmp_path):
    p = _edit(tmp_path, lambda wb: wb["4.9"].__setitem__("B18", "Universal credit (renamed)"))
    with pytest.raises(LayoutError):
        obr_efo.parse(p, _meta())


def test_universal_credit_is_the_in_cap_row_not_the_outside_cap_row():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["4.9"]
    assert ws["B43"].value == "Universal credit"  # outside the welfare cap: must not be used
    assert _by(UC)["2024-25"].value != ws["C43"].value
