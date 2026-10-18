"""Scotland GERS 2025-26 parser against a trimmed real fixture (sheets 1.1, 3.1, 3.3, 5.1)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, scotland_gers
from ukpm.provenance import problems
from ukpm.registry import FileEntry, Source

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests/fixtures/scotland_gers/gers_2025_26.trimmed.xlsx"
SHA = "61e2e8f64b280a57c61423d523c09c9252aa2f002aab7f638c129e57dcd55d9f"
URL = (
    "https://www.gov.scot/binaries/content/documents/govscot/publications/statistics/2026/08/"
    "government-expenditure-revenue-scotland-2025-26/documents/gers-2025-26-tables/"
    "gers-2025-26-tables/govscot%3Adocument/gers-2025-26-tables.xlsx"
)
LANDING = "https://www.gov.scot/publications/government-expenditure-revenue-scotland-2025-26/"


def meta(path: Path = FIXTURE) -> FileMeta:
    f = FileEntry(
        "gers_2025_26_tables", "GERS 2025-26 tables", "2026-08-12", URL, LANDING, "xlsx", SHA,
        "tables",
    )  # fmt: skip
    s = Source(
        id="scotland_gers", publisher="Scottish Government", dataset="GERS 2025-26",
        landing_url=LANDING, licence="OGL-UK-3.0", licence_evidence=None, cadence=None,
        coverage={}, verdict="USABLE", enabled=True, fixture_dir=None, user_agent=None, files=(f,),
    )  # fmt: skip
    return FileMeta(s, f, SHA, URL, "2026-10-18T00:00:00Z", path)


@pytest.fixture(scope="module")
def rows():
    return scotland_gers.parse(FIXTURE, meta())


def one(rows, sid, fy):
    (r,) = [x for x in rows if x.series_id == f"scotland_gers.{sid}" and x.fiscal_year == fy]
    return r


def edited(tmp_path, fn):
    wb = openpyxl.load_workbook(FIXTURE)
    fn(wb)
    p = tmp_path / "x.xlsx"
    wb.save(p)
    return p


def test_headline_values(rows):
    assert one(rows, "revenue", "2025-26").value == pytest.approx(98261.89)
    assert one(rows, "revenue_excl_north_sea", "2025-26").value == pytest.approx(95108.94)
    assert one(rows, "expenditure", "2025-26").value == 123566
    assert one(rows, "net_balance", "2025-26").value == pytest.approx(-25303.64)
    assert one(rows, "net_balance_excl_north_sea", "2025-26").value == pytest.approx(-28456.6)
    assert one(rows, "revenue", "1998-99").value == pytest.approx(31632.03)
    assert one(rows, "fn_health", "2025-26").value == pytest.approx(21120.07)
    assert one(rows, "fn_total_expenditure", "2025-26").value == pytest.approx(123565.53)


def test_counts_and_series_declared(rows):
    assert {r.series_id for r in rows} == set(scotland_gers.SERIES)
    for sid in ("revenue", "expenditure", "net_balance"):
        assert len([r for r in rows if r.series_id == f"scotland_gers.{sid}"]) == 28
    fn = [r for r in rows if r.series_id.startswith("scotland_gers.fn_")]
    assert {r.fiscal_year for r in fn} == {"2025-26"}
    assert all(r.geography == "Scotland" and r.unit == "gbp_m" for r in rows)


def test_function_lines_sum_to_total(rows):
    parts = [
        r.value for r in rows
        if r.series_id.startswith("scotland_gers.fn_")
        and not r.series_id.endswith("total_expenditure")
    ]  # fmt: skip
    assert sum(parts) == pytest.approx(one(rows, "fn_total_expenditure", "2025-26").value, abs=0.6)


def test_provenance_and_cells(rows):
    wb = openpyxl.load_workbook(FIXTURE, data_only=True)
    for r in rows:
        assert problems(r.prov) == [], r.series_id
        assert (r.prov.file_sha256, r.prov.source_url, r.prov.landing_url) == (SHA, URL, LANDING)
        assert r.prov.licence == "OGL-UK-3.0" and r.prov.transformations == ()
        assert wb[r.prov.sheet][r.prov.cell_or_range].value == pytest.approx(r.value)
    assert one(rows, "revenue", "2025-26").prov.sheet == "Table_1.1"
    assert one(rows, "revenue", "2025-26").prov.cell_or_range == "AC3"
    assert one(rows, "fn_total_expenditure", "2025-26").prov.cell_or_range == "B22"


def test_wrong_title_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_1.1"].__setitem__("A1", "Table 1.1: Something else"))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_missing_sheet_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb.remove(wb["Table_5.1"]))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_renamed_label_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_3.1"].__setitem__("A3", "Scotland (£ billion)"))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_year_header_change_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_5.1"].__setitem__("C2", "2000-01"))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_non_numeric_value_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_1.1"].__setitem__("AC3", "[N/A]"))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_balance_mismatch_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_5.1"].__setitem__("AC3", -20000))
    with pytest.raises(LayoutError, match="net_balance"):
        scotland_gers.parse(p, meta(p))


def test_function_year_and_unit_checked(tmp_path):
    p = edited(
        tmp_path,
        lambda wb: wb["Table_3.3"].__setitem__(
            "A1", "Table 3.3: Total Expenditure: Scotland 2024-25"
        ),
    )
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))
    p = edited(tmp_path, lambda wb: wb["Table_3.3"].__setitem__("B2", "£ billion"))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_function_total_mismatch_raises(tmp_path):
    p = edited(tmp_path, lambda wb: wb["Table_3.3"].__setitem__("B16", 22000))
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))


def test_replaced_function_line_raises(tmp_path):
    def add(wb):
        ws = wb["Table_3.3"]
        ws["A21"] = "Brand new function"
        ws["B21"] = 1

    p = edited(tmp_path, add)
    with pytest.raises(LayoutError):
        scotland_gers.parse(p, meta(p))
