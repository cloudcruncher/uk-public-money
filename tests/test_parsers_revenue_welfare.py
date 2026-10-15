"""HMRC receipts, HMRC income tax and DWP benefit parsers against their trimmed fixtures."""

from __future__ import annotations

from dataclasses import replace
from functools import cache
from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, dwp_benefits, hmrc_income_tax, hmrc_receipts
from ukpm.provenance import validate
from ukpm.registry import load

ROOT = Path(__file__).resolve().parent.parent
FX = ROOT / "tests/fixtures"
FETCHED = "2026-10-15T00:00:00Z"
PATHS = {
    "hmrc_ns_table": FX / "revenue/NS_Table.trimmed.xlsx",
    "hmrc_table_2_5": FX / "revenue/Table_2.5.trimmed.xlsx",
    "hmrc_table_2_1": FX / "revenue/Table_2.1.trimmed.xlsx",
    "dwp_spring_2026": FX / "dwp/spring_forecast_2026.trimmed.xlsx",
}


@cache
def _sources():
    return {s.id: s for s in load(ROOT / "data/sources.yml")}


def _meta(source_id, key, path=None):
    s = _sources()[source_id]
    f = next(x for x in s.files if x.key == key)
    p = path or PATHS[key]
    return FileMeta(s, f, f.sha256, f.file_url, FETCHED, p)


@cache
def _wb(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def _parse(mod, source_id, keys):
    out = []
    for k in keys:
        out += mod.parse(PATHS[k], _meta(source_id, k))
    return out


@pytest.fixture(scope="module")
def receipts():
    return _parse(hmrc_receipts, "hmrc_receipts", ["hmrc_ns_table"])


@pytest.fixture(scope="module")
def itax():
    return _parse(hmrc_income_tax, "hmrc_income_tax", ["hmrc_table_2_5", "hmrc_table_2_1"])


@pytest.fixture(scope="module")
def dwp():
    return _parse(dwp_benefits, "dwp_benefits", ["dwp_spring_2026"])


def _one(rows, series, fy):
    (r,) = [x for x in rows if x.series_id == series and x.fiscal_year == fy]
    return r


def _check_all(rows, path_of):
    assert rows
    for r in rows:
        validate(r.prov)
        assert r.prov.licence == "OGL-UK-3.0" and r.prov.transformations == ()
        cell = _wb(path_of(r))[r.prov.sheet][r.prov.cell_or_range].value
        assert cell == r.value, (r.series_id, r.fiscal_year, r.prov.cell_or_range)
    keys = [(r.series_id, r.fiscal_year) for r in rows]
    assert len(keys) == len(set(keys))


# ---- HMRC receipts -------------------------------------------------------------------------


def test_receipts_every_row_matches_its_recorded_cell(receipts):
    _check_all(receipts, lambda r: PATHS["hmrc_ns_table"])
    assert {r.prov.sheet for r in receipts} == {"Receipts_Annually"}


@pytest.mark.parametrize(
    ("fy", "value", "cell"),
    [("2015-16", 532504.0, "C16"), ("2024-25", 858648.0, "C25"), ("2025-26", 937763.0, "C26")],
)
def test_receipts_total_per_year(receipts, fy, value, cell):
    r = _one(receipts, "hmrc_receipts.total", fy)
    assert (r.value, r.prov.cell_or_range, r.unit, r.geography) == (value, cell, "gbp_m", "UK")
    assert r.prov.status == "outturn"
    assert _wb(PATHS["hmrc_ns_table"])["Receipts_Annually"][cell].value == value


def test_receipts_all_outturn_and_year_set(receipts):
    assert {r.prov.status for r in receipts} == {"outturn"}
    assert {r.fiscal_year for r in receipts} == {"2015-16", "2024-25", "2025-26"}


def test_receipts_series_slugs(receipts):
    ids = {r.series_id for r in receipts}
    for s in ("total", "total_paid_over", "income_tax", "income_tax_paye", "nics", "vat"):
        assert f"hmrc_receipts.{s}" in ids


def _edit(tmp_path, key, fn):
    w = openpyxl.load_workbook(PATHS[key])
    fn(w)
    p = tmp_path / "e.xlsx"
    w.save(p)
    return p


def test_receipts_anchor_missing_raises(tmp_path):
    def f(w):
        for row in w["Receipts_Annually"].iter_rows():
            for c in row:
                if c.value == hmrc_receipts.ANCHOR:
                    c.value = "Table 2 changed"

    p = _edit(tmp_path, "hmrc_ns_table", f)
    with pytest.raises(LayoutError, match="Table 2 HMRC"):
        hmrc_receipts.parse(p, _meta("hmrc_receipts", "hmrc_ns_table", p))


def test_receipts_sheet_missing_raises(tmp_path):
    p = _edit(tmp_path, "hmrc_ns_table", lambda w: setattr(w["Receipts_Annually"], "title", "X"))
    with pytest.raises(LayoutError):
        hmrc_receipts.parse(p, _meta("hmrc_receipts", "hmrc_ns_table", p))


def test_receipts_missing_total_column_raises(tmp_path):
    def f(w):
        for row in w["Receipts_Annually"].iter_rows():
            for c in row:
                if c.value == "Total HMRC Receipts":
                    c.value = "Grand total"

    p = _edit(tmp_path, "hmrc_ns_table", f)
    with pytest.raises(LayoutError, match="Total HMRC Receipts"):
        hmrc_receipts.parse(p, _meta("hmrc_receipts", "hmrc_ns_table", p))


def test_receipts_blank_and_text_cells_give_no_row(tmp_path, receipts):
    p = _edit(tmp_path, "hmrc_ns_table", lambda w: w["Receipts_Annually"].__setitem__("C16", "n/a"))
    out = hmrc_receipts.parse(p, _meta("hmrc_receipts", "hmrc_ns_table", p))
    assert len(out) == len(receipts) - 1
    assert not [
        r for r in out if r.series_id == "hmrc_receipts.total" and r.fiscal_year == "2015-16"
    ]


# ---- HMRC income tax -----------------------------------------------------------------------


def _itax_path(r):
    return PATHS["hmrc_table_2_1" if r.prov.sheet.startswith("2_1") else "hmrc_table_2_5"]


def test_itax_every_row_matches_its_recorded_cell(itax):
    _check_all(itax, _itax_path)
    assert {r.period_basis for r in itax} == {"tax_year"}


@pytest.mark.parametrize(
    ("fy", "value", "cell", "status"),
    [
        ("2015-16", 31000.0, "B29", "outturn"),
        ("2023-24", 36700.0, "B37", "outturn"),
        ("2024-25", 38600.0, "B38", "forecast"),
        ("2026-27", 40800.0, "B40", "forecast"),
    ],
)
def test_itax_payers_all_per_year(itax, fy, value, cell, status):
    r = _one(itax, "hmrc_income_tax.payers.all", fy)
    assert (r.value, r.prov.cell_or_range, r.prov.status, r.unit) == (
        value,
        cell,
        status,
        "thousand_persons",
    )
    assert r.prov.sheet == "2_1_IT_payers_by_demographic"
    assert _wb(PATHS["hmrc_table_2_1"])[r.prov.sheet][cell].value == value


@pytest.mark.parametrize(
    ("fy", "value", "status"), [("2023-24", 274000.0, "outturn"), ("2026-27", 347000.0, "forecast")]
)
def test_itax_liabilities_total_per_year(itax, fy, value, status):
    r = _one(itax, "hmrc_income_tax.liabilities_total", fy)
    assert (r.value, r.prov.cell_or_range, r.prov.status, r.unit) == (value, "L15", status, "gbp_m")
    assert r.prov.sheet == f"2_5_IT_Liabilities_{fy}"


def test_itax_status_boundary_and_partial_coverage(itax):
    for r in itax:
        assert r.prov.status == ("outturn" if r.fiscal_year <= "2023-24" else "forecast")
    liabs = {r.fiscal_year for r in itax if r.series_id == "hmrc_income_tax.liabilities_total"}
    assert liabs == {"2023-24", "2026-27"}  # documented partial coverage, no invented years
    assert {r.fiscal_year for r in itax if r.series_id == "hmrc_income_tax.payers.all"} == {
        "2015-16",
        "2023-24",
        "2024-25",
        "2026-27",
    }


def test_itax_band_series_per_sheet_share_label_digits(itax):
    bands = {r.series_id for r in itax if ".band.from_" in r.series_id}
    assert "hmrc_income_tax.payers.band.from_12570" in bands
    assert "hmrc_income_tax.liabilities.band.from_1000000" in bands


def test_itax_layout_errors(tmp_path):
    def kill25(w):
        for ws in w:
            for row in ws.iter_rows():
                for c in row:
                    if c.value == "Range of total income (lower limit)":
                        c.value = "Range"

    p = _edit(tmp_path, "hmrc_table_2_5", kill25)
    with pytest.raises(LayoutError, match="lower limit"):
        hmrc_income_tax.parse(p, _meta("hmrc_income_tax", "hmrc_table_2_5", p))

    p = _edit(
        tmp_path,
        "hmrc_table_2_5",
        lambda w: [setattr(ws, "title", "Z" + str(i)) for i, ws in enumerate(w)],
    )
    with pytest.raises(LayoutError, match="IT_Liabilities"):
        hmrc_income_tax.parse(p, _meta("hmrc_income_tax", "hmrc_table_2_5", p))

    p = _edit(
        tmp_path,
        "hmrc_table_2_1",
        lambda w: setattr(w["2_1_IT_payers_by_demographic"], "title", "Q"),
    )
    with pytest.raises(LayoutError, match="Tax year"):
        hmrc_income_tax.parse(p, _meta("hmrc_income_tax", "hmrc_table_2_1", p))


def test_itax_unknown_file_key_raises():
    meta = _meta("hmrc_income_tax", "hmrc_table_2_1")
    meta = replace(meta, file=replace(meta.file, key="other"))
    with pytest.raises(LayoutError, match="unknown file key"):
        hmrc_income_tax.parse(PATHS["hmrc_table_2_1"], meta)


def test_itax_no_estimate_text_gives_no_row(tmp_path, itax):
    p = _edit(
        tmp_path,
        "hmrc_table_2_1",
        lambda w: w["2_1_IT_payers_by_demographic"].__setitem__("B29", "[no estimate]"),
    )
    out = hmrc_income_tax.parse(p, _meta("hmrc_income_tax", "hmrc_table_2_1", p))
    assert not [
        r for r in out if r.series_id == "hmrc_income_tax.payers.all" and r.fiscal_year == "2015-16"
    ]


# ---- DWP benefits --------------------------------------------------------------------------

DWP_YEARS = {"2015-16", "2023-24", "2025-26", "2026-27"}


def test_dwp_every_row_matches_its_recorded_cell(dwp):
    _check_all(dwp, lambda r: PATHS["dwp_spring_2026"])
    assert {r.fiscal_year for r in dwp} == DWP_YEARS
    assert {r.prov.sheet for r in dwp} == {"Table 1a", "UK welfare "}


GB_TOTAL = {
    "2015-16": (171893.0377999211, "BS80", "outturn"),
    "2023-24": (267872.0435623201, "CA80", "outturn"),
    "2025-26": (309057.7726397607, "CC80", "forecast"),
    "2026-27": (328980.6918059892, "CD80", "forecast"),
}
UK_TOTAL = {
    "2015-16": (217.9441869480611, "X19", "outturn"),
    "2023-24": (286.963922252727, "AF19", "outturn"),
    "2025-26": (333.7771016251841, "AH19", "forecast"),
    "2026-27": (354.1515973880565, "AI19", "forecast"),
}


@pytest.mark.parametrize("fy", sorted(DWP_YEARS))
def test_dwp_totals_per_year(dwp, fy):
    v, cell, st = GB_TOTAL[fy]
    r = _one(dwp, "dwp_benefits.exp.total_benefit_expenditure", fy)
    assert (r.value, r.prov.cell_or_range, r.prov.status, r.unit, r.geography, r.prov.sheet) == (
        v,
        cell,
        st,
        "gbp_m",
        "GB",
        "Table 1a",
    )
    v, cell, st = UK_TOTAL[fy]
    r = _one(dwp, "dwp_benefits.total_uk_welfare", fy)
    assert (r.value, r.prov.cell_or_range, r.prov.status, r.unit, r.geography, r.prov.sheet) == (
        v,
        cell,
        st,
        "gbp_bn",
        "UK",
        "UK welfare ",
    )
    assert _wb(PATHS["dwp_spring_2026"])["UK welfare "][cell].value == v


def test_dwp_status_comes_from_sheet_not_from_year(dwp):
    # 2025-26 is labelled Forecast in the Spring 2026 file.
    assert {r.prov.status for r in dwp if r.fiscal_year == "2025-26"} == {"forecast"}
    assert {r.prov.status for r in dwp if r.fiscal_year in ("2015-16", "2023-24")} == {"outturn"}


def test_dwp_of_which_rows_are_children_and_obr_block_excluded(dwp):
    ids = {r.series_id for r in dwp}
    assert "dwp_benefits.exp.state_pension.contributory" in ids
    assert "dwp_benefits.exp.universal_credit.within_welfare_cap" in ids
    assert not [i for i in ids if "reconciliation" in i]
    assert "dwp_benefits.uk_welfare.total_obr_definition" in ids  # kept apart from the total


def test_dwp_layout_errors(tmp_path):
    def rn(title):
        def f(w):
            w[title].title = "gone"

        return f

    p = _edit(tmp_path, "dwp_spring_2026", rn("Table 1a"))
    with pytest.raises(LayoutError, match="Table 1a"):
        dwp_benefits.parse(p, _meta("dwp_benefits", "dwp_spring_2026", p))
    p = _edit(tmp_path, "dwp_spring_2026", rn("UK welfare "))
    with pytest.raises(LayoutError, match="UK welfare"):
        dwp_benefits.parse(p, _meta("dwp_benefits", "dwp_spring_2026", p))


def test_dwp_missing_status_label_raises(tmp_path):
    def f(w):
        ws = w["Table 1a"]
        for row in ws.iter_rows():
            for c in row:
                if c.value == "Table 1a: Expenditure by benefit," or (
                    isinstance(c.value, str) and c.value.startswith(dwp_benefits.T1A_ANCHOR)
                ):
                    hr = c.row
        for c in ws[hr + 1]:
            if c.value in ("Outturn", "Forecast"):
                c.value = "TBC"

    p = _edit(tmp_path, "dwp_spring_2026", f)
    with pytest.raises(LayoutError, match="status"):
        dwp_benefits.parse(p, _meta("dwp_benefits", "dwp_spring_2026", p))
