"""OBR Public Finances Databank parser against the trimmed fixture (cells re-opened per row)."""

from __future__ import annotations

import re
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, obr
from ukpm.provenance import validate
from ukpm.registry import load

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests/fixtures/obr/PSF_aggregates_databank_Sep-3.trimmed.xlsx"
YEARS = [f"{y}-{str(y + 1)[2:]}" for y in range(2015, 2031)]
OUTTURN = YEARS[:11]
FORECAST = YEARS[11:]


@pytest.fixture(scope="module")
def src():
    return next(s for s in load(ROOT / "data/sources.yml") if s.id == "obr")


def _meta(src, path=FIXTURE, edition=None):
    f = src.files[0]
    if edition:
        f = replace(f, edition=edition)
    return FileMeta(src, f, f.sha256, f.file_url, "2026-10-15T00:00:00Z", path)


@pytest.fixture(scope="module")
def rows(src):
    return obr.parse(FIXTURE, _meta(src))


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(FIXTURE, data_only=True)


def _one(rows, series, fy):
    (r,) = [x for x in rows if x.series_id == series and x.fiscal_year == fy]
    return r


def test_every_row_has_complete_provenance_and_value_equals_recorded_cell(rows, wb):
    assert len(rows) > 500
    for r in rows:
        validate(r.prov)
        assert r.prov.vintage == "OBR databank Sep 2026"
        assert r.prov.licence == "OGL-UK-3.0"
        assert r.prov.transformations == ()
        assert r.geography == "UK" and r.period_basis == "fiscal_year"
        cell = wb[r.prov.sheet][r.prov.cell_or_range].value
        assert cell == pytest.approx(r.value, rel=1e-12), (r.series_id, r.fiscal_year)


def test_no_duplicate_series_year_and_no_year_before_2015_16(rows):
    keys = [(r.series_id, r.fiscal_year) for r in rows]
    assert len(keys) == len(set(keys))
    assert min(r.fiscal_year for r in rows) == "2015-16"
    assert max(r.fiscal_year for r in rows) == "2030-31"


TME = [
    794.365, 812.735, 842.25, 857.627, 886.838, 1103.793, 1038.195, 1158.392,
    1231.793, 1290.348, 1365.598,
]  # fmt: skip
TME_FORECAST = [1419.254875556987, 1471.874916353111, 1513.453239031104, 1555.347858666221,
                1610.504465127791]  # fmt: skip


@pytest.mark.parametrize("i", range(16))
def test_tme_per_year_value_cell_and_status(rows, wb, i):
    fy = YEARS[i]
    r = _one(rows, "obr.psf.tme", fy)
    want = (TME + TME_FORECAST)[i]
    assert r.value == pytest.approx(want, rel=1e-12)
    assert (r.prov.sheet, r.prov.cell_or_range) == (obr.AGG, f"D{77 + i}")
    assert r.prov.status == ("outturn" if fy in OUTTURN else "forecast")
    assert r.unit == "gbp_bn"


def test_status_boundary_is_2026_27_for_every_series(rows):
    assert rows
    for r in rows:
        assert r.prov.status == ("forecast" if r.fiscal_year >= "2026-27" else "outturn"), r


def test_core_series_all_present_for_all_sixteen_years(rows):
    for sid in (
        "obr.psf.receipts",
        "obr.psf.tme",
        "obr.psf.current_expenditure",
        "obr.psf.net_investment",
        "obr.psf.depreciation",
        "obr.psf.gross_investment",
        "obr.gdp_nominal",
        "obr.gdp_deflator",
    ):
        assert sorted(r.fiscal_year for r in rows if r.series_id == sid) == YEARS, sid


def test_gdp_and_deflator_known_values_and_units(rows):
    g = _one(rows, "obr.gdp_nominal", "2025-26")
    assert (g.value, g.prov.sheet, g.prov.cell_or_range, g.unit) == (
        3069.799,
        obr.PCT,
        "AE82",
        "gbp_bn",
    )
    d = _one(rows, "obr.gdp_deflator", "2025-26")
    assert (d.value, d.prov.sheet, d.prov.cell_or_range, d.unit) == (
        100.0,
        obr.REAL,
        "AE75",
        "index",
    )
    assert _one(rows, "obr.gdp_deflator", "2015-16").value == pytest.approx(72.09302325581395)
    assert _one(rows, "obr.gdp_nominal", "2030-31").prov.status == "forecast"


def test_receipts_by_tax_series_from_receipts_sheet(rows):
    vat = [r for r in rows if r.series_id == "obr.receipts.vat"]
    assert sorted(r.fiscal_year for r in vat) == YEARS
    assert {r.prov.sheet for r in vat} == {obr.REC}
    # skipped columns are not duplicated from the receipts sheet
    assert not [r for r in rows if r.prov.sheet == obr.REC and r.series_id == "obr.psf.receipts"]
    assert "obr.receipts.paye_income_tax" in {r.series_id for r in rows}
    assert "obr.receipts.nics" in {r.series_id for r in rows}


def test_series_registered_for_dynamic_receipts(rows):
    assert "obr.receipts.vat" in obr.SERIES
    assert obr.SERIES["obr.receipts.vat"]["unit"] == "gbp_bn"


# ---- layout handling -----------------------------------------------------------------------


def _edit(tmp_path, fn):
    w = openpyxl.load_workbook(FIXTURE)
    fn(w)
    p = tmp_path / "edited.xlsx"
    w.save(p)
    return p


def _find(ws, text):
    return next(c for row in ws.iter_rows() for c in row if c.value == text)


def test_missing_aggregates_anchor_raises(src, tmp_path):
    p = _edit(
        tmp_path, lambda w: setattr(_find(w[obr.AGG], "Total managed expenditure"), "value", "TME")
    )
    with pytest.raises(LayoutError, match="Total managed expenditure"):
        obr.parse(p, _meta(src, p))


def test_missing_sheet_raises_layout_error_or_key_error_not_silent(src, tmp_path):
    p = _edit(tmp_path, lambda w: setattr(w[obr.REC], "title", "Receipts"))
    with pytest.raises((LayoutError, KeyError)):
        obr.parse(p, _meta(src, p))


def test_missing_forecast_note_raises(src, tmp_path):
    def kill(w):
        for ws in (w[obr.AGG],):
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and "Forecast years from" in c.value:
                        c.value = "note removed"

    p = _edit(tmp_path, kill)
    with pytest.raises(LayoutError, match="Forecast years"):
        obr.parse(p, _meta(src, p))


def test_forecast_boundary_follows_the_sheet_note(src, tmp_path):
    def shift(w):
        for row in w[obr.AGG].iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Forecast years from 2026-27" in c.value:
                    c.value = c.value.replace("2026-27", "2028-29")

    p = _edit(tmp_path, shift)
    out = obr.parse(p, _meta(src, p))
    tme = {r.fiscal_year: r.prov.status for r in out if r.series_id == "obr.psf.tme"}
    assert tme["2027-28"] == "outturn" and tme["2028-29"] == "forecast"


def test_missing_deflator_column_raises(src, tmp_path):
    def kill(w):
        for row in w[obr.REAL].iter_rows(max_row=15):
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("GDP Deflator"):
                    c.value = "x"

    p = _edit(tmp_path, kill)
    with pytest.raises(LayoutError, match="Deflator"):
        obr.parse(p, _meta(src, p))


def test_dash_cell_gives_no_row(src, tmp_path):
    p = _edit(tmp_path, lambda w: w[obr.AGG].__setitem__("D77", "-"))
    out = obr.parse(p, _meta(src, p))
    assert not [r for r in out if r.series_id == "obr.psf.tme" and r.fiscal_year == "2015-16"]
    assert len([r for r in out if r.series_id == "obr.psf.tme"]) == 15


def test_vintage_derived_from_edition_and_bad_edition_rejected(src):
    assert obr._vintage("OBR Public Finances Databank Mar 2027") == "OBR databank Mar 2027"
    assert obr._vintage("OBR Databank September 2026") == "OBR databank Sep 2026"
    with pytest.raises(ValueError):
        obr.parse(FIXTURE, _meta(src, edition="OBR databank latest"))


@pytest.mark.parametrize(
    ("label", "slug"),
    [
        ("Pay as your earn income tax", "paye_income_tax"),
        ("Self assessed income tax", "sa_income_tax"),
        ("National insurance contributions", "nics"),
        ("Stamp duty land tax1", "stamp_duty_land_tax"),
        ("VAT (net of VAT refunds)", "vat"),
    ],
)
def test_slug(label, slug):
    assert obr._slug(label) == slug
    assert re.fullmatch(r"[a-z0-9_]+", slug)


def test_databank_real_terms_tme_extracted_as_published(rows):
    assert sorted(r.fiscal_year for r in rows if r.series_id == "obr.psf.tme_real") == YEARS
    r = _one(rows, "obr.psf.tme_real", "2025-26")
    # base year of the deflator (=100): real equals nominal
    assert r.value == pytest.approx(_one(rows, "obr.psf.tme", "2025-26").value)
    assert (r.prov.sheet, r.unit, r.prov.transformations) == (obr.REAL, "gbp_bn", ())
    assert r.prov.status == "outturn"
    assert _one(rows, "obr.psf.tme_real", "2027-28").prov.status == "forecast"
