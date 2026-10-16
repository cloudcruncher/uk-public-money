"""Wave-2 parsers against trimmed real fixtures: ONS PSF, households, GDP, PESA 2026 chapter 5,
HMRC receipts (NS_Table) and OBR EFO receipts (table 3.8)."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import (
    FileMeta,
    LayoutError,
    hmrc_receipts,
    obr_efo_receipts,
    ons_gdp,
    ons_households,
    ons_psf,
    pesa_uk,
)
from ukpm.provenance import problems
from ukpm.registry import FileEntry, Source

ROOT = Path(__file__).resolve().parent.parent
FX = ROOT / "tests/fixtures"
ONS = "https://www.ons.gov.uk/"
GOV = "https://assets.publishing.service.gov.uk/media/"

# name -> (module, source id, fixture, original sha, file url, landing, publication date, edition)
CASES = {
    "psf": (
        ons_psf, "ons_psf", FX / "ons2/psf_aug2026.trimmed.xlsx",
        "4cc5dded99d4bf06b8df9c3d9ef4e09a903c96208ba54626e6db507a07129b7e",
        ONS + "file?uri=/economy/governmentpublicsectorandtaxes/publicsectorfinance/datasets/"
        "publicsectorfinancesappendixatables110/current/"
        "publicsectorfinancessummarytablesappendixafinal.xlsx",
        ONS + "economy/governmentpublicsectorandtaxes/publicsectorfinance/datasets/"
        "publicsectorfinancesappendixatables110",
        "2026-09-22", "ONS Public sector finances, August 2026",
    ),
    "hh": (
        ons_households, "ons_households", FX / "ons2/familiesandhouseholdsuk2025.trimmed.xlsx",
        "a89eaa3d7f519bfa949c7cfa12f531b8c7a86c704adcb70e81f49567ae724a91",
        ONS + "file?uri=/peoplepopulationandcommunity/birthsdeathsandmarriages/families/datasets/"
        "familiesandhouseholdsfamiliesandhouseholds/current/familiesandhouseholdsuk2025.xlsx",
        ONS + "peoplepopulationandcommunity/birthsdeathsandmarriages/families/datasets/"
        "familiesandhouseholdsfamiliesandhouseholds",
        "2026-04-17", "Families and households in the UK: 2025",
    ),
    "gdp": (
        ons_gdp, "ons_gdp", FX / "ons2/ybha_ukea.trimmed.csv",
        "080b6afad67e567c2f27c1afea720cdf49832b7158de0487fc4a3681c0485bf4",
        ONS + "generator?format=csv&uri=/economy/grossdomesticproductgdp/timeseries/ybha/ukea",
        ONS + "economy/grossdomesticproductgdp/timeseries/ybha/ukea",
        "2026-09-30", "UKEA YBHA 30 September 2026",
    ),
    "pesa": (
        pesa_uk, "pesa_uk", FX / "pesa_uk/pesa_2026_ch5.trimmed.xlsx",
        "ccb4539bb837c2b6a7f9509da49bd409eb9e5a910f3332311b494291fbd33107",
        GOV + "6a5772f59e63154454413701/PESA_2026_CP_Chapter_5_tables.xlsx",
        "https://www.gov.uk/government/statistics/public-expenditure-statistical-analyses-2026",
        "2026-07-16", "PESA 2026",
    ),
    "hmrc": (
        hmrc_receipts, "hmrc_receipts", FX / "hmrc/NS_Table.trimmed.xlsx",
        "2cadbec4c57e8626dae26c9f46975375eb2d0c07e40a94070a3c6ccc64403cab",
        GOV + "6aa923f20420ac660805f7f2/NS_Table.ods",
        "https://www.gov.uk/government/statistics/hmrc-tax-and-nics-receipts-for-the-uk",
        "2026-09-22", "HMRC receipts, September 2026 update",
    ),
    "efo": (
        obr_efo_receipts, "obr_efo_receipts",
        FX / "obr_receipts/efo_march_2026_receipts.trimmed.xlsx",
        "f65b6cb7f96931d4abd85940794eb60f6f65afca8bbf16c74a4548bf92766a28",
        "https://obr.uk/docs/d055fbf02d5b3g6jq8l2/"
        "efo-march-2026-detailed-forecast-tables-receipts.xlsx",
        "https://obr.uk/efo/economic-and-fiscal-outlook-march-2026/",
        "2026-03-26", "OBR EFO March 2026",
    ),
}  # fmt: skip


def meta_for(name: str, path: Path | None = None) -> FileMeta:
    mod, sid, fixture, sha, url, landing, pub, edition = CASES[name]
    f = FileEntry(f"{sid}_k", edition, pub, url, landing, "xlsx", sha, "main")
    s = Source(
        id=sid, publisher="Test publisher", dataset=f"{sid} dataset", landing_url=landing,
        licence="OGL-UK-3.0", licence_evidence=None, cadence=None, coverage={},
        verdict="USABLE", enabled=True, fixture_dir=None, user_agent=None, files=(f,),
    )  # fmt: skip
    return FileMeta(s, f, sha, url, "2026-10-15T00:00:00Z", path or fixture)


def run(name: str, path: Path | None = None):
    return CASES[name][0].parse(path or CASES[name][2], meta_for(name, path))


@pytest.fixture(scope="module")
def all_rows():
    return {n: run(n) for n in CASES}


def one(rows, sid, fy):
    (r,) = [x for x in rows if x.series_id == sid and x.fiscal_year == fy]
    return r


def edited(tmp_path, name, fn):
    wb = openpyxl.load_workbook(CASES[name][2])
    fn(wb)
    p = tmp_path / "x.xlsx"
    wb.save(p)
    return p


# ---------- provenance, all parsers ----------


@pytest.mark.parametrize("name", list(CASES))
def test_provenance_complete_and_series_declared(all_rows, name):
    rows = all_rows[name]
    assert rows
    mod, _sid, _fx, sha, url, landing, pub, _ed = CASES[name]
    for r in rows:
        assert problems(r.prov) == [], (r.series_id, r.fiscal_year)
        assert (r.prov.file_sha256, r.prov.source_url, r.prov.landing_url) == (sha, url, landing)
        assert r.prov.licence == "OGL-UK-3.0" and r.prov.transformations == ()
        assert r.prov.publication_date == pub and r.prov.vintage and r.prov.edition
        assert name == "hmrc" or r.series_id in mod.SERIES  # hmrc series are per column
    assert name == "hmrc" or {r.series_id for r in rows} == set(mod.SERIES)
    for s in mod.SERIES.values():
        assert {"title", "unit", "period_type", "monetary", "description"} <= set(s)


@pytest.mark.parametrize("name", [n for n in CASES if CASES[n][2].suffix == ".xlsx"])
def test_cell_matches_workbook(all_rows, name):
    wb = openpyxl.load_workbook(CASES[name][2], data_only=True)
    for r in all_rows[name]:
        assert wb[r.prov.sheet][r.prov.cell_or_range].value == pytest.approx(r.value), r.series_id


# ---------- ONS PSF ----------


def test_psf_tme_and_gross_expenditure_are_separate_definitions(all_rows):
    rows = all_rows["psf"]
    tme = one(rows, "ons_psf.cg_tme", "2025-26")
    gross = one(rows, "ons_psf.cg_total_expenditure_gross", "2025-26")
    assert (tme.value, tme.prov.sheet, tme.prov.cell_or_range) == (1250068.0, "PSA6L", "B63")
    assert (gross.value, gross.prov.sheet) == (1301598.0, "PSA6C_3")
    assert gross.prov.cell_or_range == "C63"
    assert gross.value - tme.value == 51530
    t = ons_psf.SERIES
    assert t["ons_psf.cg_tme"]["title"] != t["ons_psf.cg_total_expenditure_gross"]["title"]
    assert "never combine" in t["ons_psf.cg_total_expenditure_gross"]["description"]
    assert one(rows, "ons_psf.cg_tme", "2024-25").value == 1187834.0


@pytest.mark.parametrize(
    ("sid", "fy", "value", "sheet", "cell", "unit"),
    [
        ("cg_receipts", "2025-26", 1175063.0, "PSA6C_3", "B63", "gbp_m"),
        ("cg_receipts", "2024-25", 1090205.0, "PSA6C_3", "B62", "gbp_m"),
        ("cg_net_borrowing", "2025-26", 126535.0, "PSA6C_3", "F63", "gbp_m"),
        ("cg_net_borrowing", "2024-25", 150033.0, "PSA6C_3", "F62", "gbp_m"),
        ("cg_debt_interest", "2025-26", 96839.0, "PSA6B_2", "B63", "gbp_m"),
        ("cg_debt_interest", "2024-25", 85261.0, "PSA6B_2", "B62", "gbp_m"),
        ("ps_net_borrowing_ex_banks", "2025-26", 134286.0, "PSA1", "D63", "gbp_m"),
        ("ps_net_debt_gbp_bn", "2025-26", 2916.8, "PSA1", "G63", "gbp_bn"),
        ("ps_net_debt_pct_gdp", "2025-26", 93.3, "PSA1", "H63", "pct_gdp"),
    ],
)
def test_psf_values(all_rows, sid, fy, value, sheet, cell, unit):
    r = one(all_rows["psf"], f"ons_psf.{sid}", fy)
    assert (r.value, r.prov.sheet, r.prov.cell_or_range, r.unit) == (value, sheet, cell, unit)
    assert (r.period_basis, r.geography, r.prov.status) == ("fiscal_year", "UK", "outturn")
    assert r.prov.vintage == "ONS Public sector finances, published 22 September 2026"


def test_psf_years_fy_only_from_2015_16(all_rows):
    years = sorted({r.fiscal_year for r in all_rows["psf"]})
    assert years[0] == "2015-16" and years[-1] == "2025-26" and len(years) == 11


def test_psf_cdid_with_sign_prefix_and_space_found(all_rows):
    # ' -NMFJ' and ' -J5II' carry a space and a sign in row 6; located by stripped CDID
    assert one(all_rows["psf"], "ons_psf.cg_net_borrowing", "2025-26").value == 126535.0


def test_psf_missing_cdid_raises(tmp_path):
    p = edited(tmp_path, "psf", lambda w: w["PSA6L"].__setitem__("B6", "ZZZZ"))
    with pytest.raises(LayoutError):
        run("psf", p)


def test_psf_wrong_unit_heading_raises(tmp_path):
    p = edited(tmp_path, "psf", lambda w: w["PSA1"].__setitem__("G5", "Net debt (£ million)"))
    with pytest.raises(LayoutError):
        run("psf", p)


def test_psf_tme_vs_memo_total_mismatch_raises(tmp_path):
    p = edited(tmp_path, "psf", lambda w: w["PSA6B_2"].__setitem__("K63", 1.0))
    with pytest.raises(LayoutError):
        run("psf", p)


def test_psf_missing_sheet_raises(tmp_path):
    p = edited(tmp_path, "psf", lambda w: w["PSA6C_3"].__setattr__("title", "gone"))
    with pytest.raises(LayoutError):
        run("psf", p)


def test_psf_no_cover_date_raises(tmp_path):
    p = edited(tmp_path, "psf", lambda w: w["Cover Sheet"].__setitem__("A5", "no date"))
    with pytest.raises(LayoutError):
        run("psf", p)


# ---------- ONS households ----------


@pytest.mark.parametrize(
    ("fy", "value", "cell"),
    [("2025-26", 29003.0, "B21"), ("2024-25", 28609.0, "E21")],
)
def test_households_values(all_rows, fy, value, cell):
    rows = all_rows["hh"]
    r = one(rows, "ons_households.uk", fy)
    assert (r.value, r.prov.sheet, r.prov.cell_or_range) == (value, "5", cell)
    assert (r.unit, r.period_basis, r.geography, r.prov.status) == (
        "thousands", "apr_jun_quarter", "UK", "outturn",
    )  # fmt: skip


def test_households_definition_break_documented():
    d = ons_households.SERIES["ons_households.uk"]["description"]
    assert "2011" in d and "note 9" in d and "sample" in d.lower()
    assert ons_households.FIRST_YEAR == 2011


def test_households_pre_2011_not_emitted(tmp_path):
    def add(wb):
        ws = wb["5"]
        ws["N12"], ws["N21"] = "2010 Estimate", 26240.0
        ws["O12"], ws["P12"] = "2010 CV", "2010 CI+/-"
        ws["Q12"], ws["Q21"] = "2011 Estimate", 26409.0

    out = run("hh", edited(tmp_path, "hh", add))
    assert not [r for r in out if r.fiscal_year == "2010-11" and r.value == 26240.0]
    assert one(out, "ons_households.uk", "2011-12").value == 26409.0


def test_households_missing_row_raises(tmp_path):
    p = edited(tmp_path, "hh", lambda w: w["5"].__setitem__("A21", "Everyone"))
    with pytest.raises(LayoutError):
        run("hh", p)


def test_households_no_cover_date_raises(tmp_path):
    p = edited(tmp_path, "hh", lambda w: w["Cover sheet"].__setitem__("A2", "x"))
    with pytest.raises(LayoutError):
        run("hh", p)


# ---------- ONS GDP ----------


@pytest.mark.parametrize(
    ("year", "value", "line"),
    [("2025", 3044594.0, 13), ("2024", 2904950.0, 12), ("2023", 2766446.0, 11)],
)
def test_gdp_values(all_rows, year, value, line):
    r = one(all_rows["gdp"], "ons_gdp.nominal", year)
    assert (r.value, r.unit, r.period_basis, r.geography) == (value, "gbp_m", "calendar_year", "UK")
    assert (r.prov.sheet, r.prov.cell_or_range) == ("ybha_ukea.trimmed.csv", f"line {line}")
    assert r.prov.status == "outturn"
    assert r.prov.vintage == "ONS UKEA release 2026-09-30"
    assert "grossdomesticproductgdp/timeseries/ybha/ukea" in r.prov.api_query


def test_gdp_only_annual_from_2015(all_rows):
    years = sorted(r.fiscal_year for r in all_rows["gdp"])
    assert years == ["2015", "2023", "2024", "2025"]


@pytest.mark.parametrize(
    ("old", "new"),
    [('"CDID","YBHA"', '"CDID","ABMI"'), ('"Unit","m"', '"Unit","bn"'),
     ('"Release date","30-09-2026"', '"Release date","soon"')],
)  # fmt: skip
def test_gdp_layout_errors(tmp_path, old, new):
    p = tmp_path / "g.csv"
    p.write_text(CASES["gdp"][2].read_text().replace(old, new))
    with pytest.raises(LayoutError):
        run("gdp", p)


def test_gdp_non_numeric_value_is_a_layout_error(tmp_path):
    p = tmp_path / "g.csv"
    p.write_text(CASES["gdp"][2].read_text().replace('"2025","3044594"', '"2025","n/a"'))
    with pytest.raises(LayoutError, match="not a number"):
        run("gdp", p)


# ---------- PESA 2026 ch5 ----------


@pytest.mark.parametrize(
    ("slug", "value", "cell"),
    [
        ("tme", 1360122.0, "G133"),
        ("services", 1227648.0, "G131"),
        ("accounting_adjustments", 132474.0, "G132"),
        ("health", 257542.0, "G79"),
        ("social_protection", 407270.0, "G121"),
        ("defence", 65418.0, "G27"),
        ("general_public_services", 165504.0, "G20"),
        ("public_debt_transactions", 130305.0, "G14"),
        ("eu_transactions", -166.0, "G130"),
    ],
)
def test_pesa_2025_26_values(all_rows, slug, value, cell):
    r = one(all_rows["pesa"], f"pesa_uk.{slug}", "2025-26")
    assert (r.value, r.prov.sheet, r.prov.cell_or_range) == (value, "5_2", cell)
    assert (r.unit, r.geography, r.prov.status, r.prov.vintage) == (
        "gbp_m", "UK", "outturn", "PESA 2026",
    )  # fmt: skip


def test_pesa_2021_22_and_span(all_rows):
    rows = all_rows["pesa"]
    assert one(rows, "pesa_uk.tme", "2021-22").value == 1040888.0
    assert one(rows, "pesa_uk.social_protection", "2021-22").value == 300292.0
    assert sorted({r.fiscal_year for r in rows}) == [
        "2021-22", "2022-23", "2023-24", "2024-25", "2025-26",
    ]  # fmt: skip


def test_pesa_totals_reconcile(all_rows):
    rows = all_rows["pesa"]
    for fy in ("2021-22", "2025-26"):
        v = {s.split(".")[1]: one(rows, s, fy).value for s in pesa_uk.SERIES}
        assert v["services"] + v["accounting_adjustments"] == v["tme"]


def test_pesa_bad_total_raises(tmp_path):
    p = edited(tmp_path, "pesa", lambda w: w["5_2"].__setitem__("G133", 1.0))
    with pytest.raises(LayoutError):
        run("pesa", p)


def test_pesa_missing_label_raises(tmp_path):
    p = edited(tmp_path, "pesa", lambda w: w["5_2"].__setitem__("B79", "Total wellbeing"))
    with pytest.raises(LayoutError):
        run("pesa", p)


def test_pesa_wrong_unit_raises(tmp_path):
    p = edited(tmp_path, "pesa", lambda w: w["5_2"].__setitem__("G3", "£ billion"))
    with pytest.raises(LayoutError):
        run("pesa", p)


# ---------- HMRC receipts (NS_Table) ----------


@pytest.mark.parametrize(
    ("slug", "value"),
    [("income_tax", 326932.0), ("nics", 200829.0), ("vat", 179861.0), ("total", 937763.0),
     ("total_paid_over", 919111.0), ("income_tax_paye", 278237.0)],
)  # fmt: skip
def test_hmrc_2025_26_values(all_rows, slug, value):
    r = one(all_rows["hmrc"], f"hmrc_receipts.{slug}", "2025-26")
    assert (r.value, r.unit, r.prov.status) == (value, "gbp_m", "outturn")
    assert r.prov.sheet == "Receipts_Annually"


def test_hmrc_corporation_tax_cell(all_rows):
    r = one(all_rows["hmrc"], "hmrc_receipts.corporation_tax", "2025-26")
    assert r.value == 95105.0
    ws = openpyxl.load_workbook(CASES["hmrc"][2])["Receipts_Annually"]
    assert ws[r.prov.cell_or_range].value == 95105.0


# ---------- OBR EFO receipts 3.8 ----------


@pytest.mark.parametrize(
    ("slug", "fy", "value", "cell", "status"),
    [
        ("income_tax", "2024-25", 302.799, "D7", "outturn"),
        ("income_tax", "2025-26", 329.2977917811605, "E7", "forecast"),
        ("income_tax", "2030-31", 428.77176433752544, "J7", "forecast"),
        ("nics", "2025-26", 200.0813950827125, "E12", "forecast"),
        ("vat", "2025-26", 180.16893868541794, "E13", "forecast"),
        ("total_hmrc", "2024-25", 858.621786848297, "D51", "outturn"),
        ("total_hmrc", "2025-26", 936.8774655415363, "E51", "forecast"),
        ("current_receipts", "2024-25", 1139.044, "D64", "outturn"),
    ],
)
def test_efo_receipts_values(all_rows, slug, fy, value, cell, status):
    r = one(all_rows["efo"], f"obr_efo_receipts.{slug}", fy)
    assert (r.prov.sheet, r.prov.cell_or_range) == ("3.8", cell)
    assert r.value == pytest.approx(value, rel=1e-12)
    assert (r.unit, r.prov.status, r.prov.vintage) == ("gbp_bn", status, "OBR EFO March 2026")


def test_efo_footnote_digit_stripped_and_nic_not_confused(all_rows):
    # 'Income tax (gross of tax credits)1' has a footnote digit; 'Other income tax' is not it
    assert one(all_rows["efo"], "obr_efo_receipts.income_tax", "2025-26").prov.cell_or_range == "E7"


def test_efo_missing_label_raises(tmp_path):
    p = edited(tmp_path, "efo", lambda w: w["3.8"].__setitem__("B13", "Sales tax"))
    with pytest.raises(LayoutError):
        run("efo", p)


def test_efo_wrong_unit_raises(tmp_path):
    p = edited(tmp_path, "efo", lambda w: w["3.8"].__setitem__("D3", "£ million"))
    with pytest.raises(LayoutError):
        run("efo", p)


def test_efo_no_status_marker_raises(tmp_path):
    def f(w):
        w["3.8"]["D4"] = None
        w["3.8"]["E4"] = None

    with pytest.raises(LayoutError):
        run("efo", edited(tmp_path, "efo", f))
