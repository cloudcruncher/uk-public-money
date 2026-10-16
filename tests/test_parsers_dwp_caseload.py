"""DWP Table 1c caseloads and the BENEFIT_PAIRS table, against the trimmed Spring Forecast 2026."""

from __future__ import annotations

from functools import cache
from pathlib import Path

import openpyxl
import pytest

from ukpm.parsers import FileMeta, LayoutError, dwp_benefits
from ukpm.provenance import problems
from ukpm.registry import load

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "tests/fixtures/dwp/spring_forecast_2026.trimmed.xlsx"
FETCHED = "2026-10-15T00:00:00Z"


@cache
def _meta(path: Path = PATH) -> FileMeta:
    s = next(x for x in load(ROOT / "data/sources.yml") if x.id == "dwp_benefits")
    f = next(x for x in s.files if x.key == "dwp_spring_2026")
    return FileMeta(s, f, f.sha256, f.file_url, FETCHED, path)


@pytest.fixture(scope="module")
def rows():
    return dwp_benefits.parse(PATH, _meta())


@pytest.fixture(scope="module")
def caseload(rows):
    return [r for r in rows if r.series_id.startswith("dwp_benefits.caseload.")]


def _one(rows, sid, fy):
    (r,) = [x for x in rows if x.series_id == sid and x.fiscal_year == fy]
    return r


def test_carers_allowance_in_payment_2025_26(caseload):
    r = _one(caseload, "dwp_benefits.caseload.carer_s_allowance.in_payment", "2025-26")
    assert (r.value, r.prov.sheet, r.prov.cell_or_range, r.prov.status) == (
        1018.0,
        "Table 1c",
        "CC12",
        "forecast",
    )
    assert (r.unit, r.geography, r.period_basis) == ("thousand_cases", "GB", "fiscal_year")


def test_pip_in_payment_2025_26_and_outturn_year(caseload):
    r = _one(caseload, "dwp_benefits.caseload.personal_independence_payment.in_payment", "2025-26")
    assert (r.value, r.prov.cell_or_range) == (3836.0, "CC52")
    r = _one(caseload, "dwp_benefits.caseload.personal_independence_payment.in_payment", "2023-24")
    assert (r.value, r.prov.cell_or_range, r.prov.status) == (3163.0, "CA52", "outturn")


def test_values_match_workbook_cells(caseload):
    ws = openpyxl.load_workbook(PATH, data_only=True)["Table 1c"]
    assert caseload
    for r in caseload:
        assert r.prov.sheet == "Table 1c"
        assert ws[r.prov.cell_or_range].value == r.value


def test_every_benefit_pair_series_exists(rows):
    ids = {r.series_id for r in rows}
    assert len(dwp_benefits.BENEFIT_PAIRS) == 16
    for p in dwp_benefits.BENEFIT_PAIRS:
        assert p["spend"] in ids, p["slug"]
        assert p["caseload"] in ids, p["slug"]
        assert p["basis"] in ("in_payment", "all_cases")
        assert p["notes"] and all(isinstance(n, str) and n for n in p["notes"])
    slugs = [p["slug"] for p in dwp_benefits.BENEFIT_PAIRS]
    assert len(set(slugs)) == 16 and all(s == s.lower() and " " not in s for s in slugs)


def test_in_payment_rule(rows):
    for p in dwp_benefits.BENEFIT_PAIRS:
        assert p["caseload"].endswith(".in_payment") == (p["basis"] == "in_payment")


def test_provenance_complete(caseload):
    assert caseload
    for r in caseload:
        assert problems(r.prov) == []
        assert r.prov.licence == "OGL-UK-3.0" and r.prov.transformations == ()
        assert r.prov.file_sha256 == _meta().sha256


def test_no_total_caseload(rows):
    assert not [r for r in rows if "total" in r.series_id and ".caseload." in r.series_id]


def test_devolution_notes_are_quoted(rows):
    by = {p["slug"]: p["notes"] for p in dwp_benefits.BENEFIT_PAIRS}
    assert any("3rd September 2018" in n for n in by["carers-allowance"])
    assert any("1st April 2020" in n for n in by["personal-independence-payment"])
    assert any("April 2024" in n for n in by["winter-fuel-payments"])
    assert any("rounded to the nearest 1,000" in n for n in by["universal-credit"])


def test_missing_1c_anchor_raises(tmp_path):
    wb = openpyxl.load_workbook(PATH)
    wb["Table 1c"]["B2"] = "Something else"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(LayoutError):
        dwp_benefits.parse(p, _meta(p))


def test_missing_1c_footnote_raises(tmp_path):
    wb = openpyxl.load_workbook(PATH)
    wb["Table 1c"]["B63"] = "gone"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(LayoutError):
        dwp_benefits.parse(p, _meta(p))
